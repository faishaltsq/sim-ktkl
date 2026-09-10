import io
import datetime
from django.test import TestCase, Client
from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
import openpyxl

from .models import (
    Nakes, AuditLog, Profesi, DokumenNakes, DokumenUmum,
    EvaluasiOPPE, EvaluasiMutuKlinis,
    PelanggaranEtik, SidangEtik, EvaluasiKinerjaEtik,
    AgendaRapat, Regulasi, NotulenRapat,
)


class ExcelImportQATest(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username='tester', password='password123')
        self.client = Client()
        self.client.login(username='tester', password='password123')
        self.p_atlm, _ = Profesi.objects.get_or_create(nama='ATLM')
        self.p_rad, _ = Profesi.objects.get_or_create(nama='Radiografer')
        self.p_fis, _ = Profesi.objects.get_or_create(nama='Fisioterapis')
        self.p_nut, _ = Profesi.objects.get_or_create(nama='Nutrisionis')
        self.p_apo, _ = Profesi.objects.get_or_create(nama='Apoteker')

    def _create_excel_file(self, rows, sheet_name='Sheet1', headers=None):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet_name
        if headers:
            ws.append(headers)
        for r in rows:
            ws.append(r)
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return SimpleUploadedFile('test.xlsx', buf.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    def test_download_template(self):
        url = reverse('nakes:download_template')
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        self.assertIn('Template Import', wb.sheetnames)
        self.assertIn('Panduan', wb.sheetnames)
        self.assertEqual(wb.active.title, 'Template Import')

    def test_upload_with_active_sheet_panduan(self):
        wb = openpyxl.Workbook()
        ws_data = wb.active
        ws_data.title = 'Template Import'
        ws_data.append(['nama', 'profesi', 'unit_kerja', 'no_str', 'masa_berlaku_str', 'no_sip', 'masa_berlaku_sip', 'status_kredensial', 'kewenangan_klinis', 'catatan'])
        ws_data.append(['Budi Panduan', 'ATLM', 'Lab', 'STR-P1', '2027-01-01', 'SIP-P1', '2027-01-01', 'Selesai', 'Aktif', ''])

        ws_guide = wb.create_sheet('Panduan')
        ws_guide['A1'] = 'PANDUAN IMPORT'
        wb.active = ws_guide

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        file = SimpleUploadedFile('import.xlsx', buf.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

        response = self.client.post(reverse('nakes:upload_bulk'), {'excel_file': file}, follow=True)
        self.assertEqual(response.status_code, 200)
        self.assertTrue(Nakes.objects.filter(no_str='STR-P1').exists())

    def test_upload_indonesian_date_formats(self):
        headers = ['nama', 'profesi', 'unit_kerja', 'no_str', 'masa_berlaku_str', 'no_sip', 'masa_berlaku_sip', 'status_kredensial', 'kewenangan_klinis', 'catatan']
        rows = [
            ['Nakes Slash', 'ATLM', 'Lab', 'STR-DT-1', '31/12/2026', 'SIP-DT-1', '31/12/2026', 'Selesai', 'Aktif', ''],
            ['Nakes Dash', 'Radiografer', 'Radiologi', 'STR-DT-2', '15-08-2027', 'SIP-DT-2', '15-08-2027', 'Selesai', 'Aktif', ''],
            ['Nakes Serial', 'Apoteker', 'Farmasi', 'STR-DT-3', 46385, 'SIP-DT-3', 46385, 'Selesai', 'Aktif', ''],
        ]
        file = self._create_excel_file(rows, headers=headers)
        response = self.client.post(reverse('nakes:upload_bulk'), {'excel_file': file}, follow=True)
        self.assertEqual(response.status_code, 200)

        n1 = Nakes.objects.get(no_str='STR-DT-1')
        self.assertEqual(n1.masa_berlaku_str, datetime.date(2026, 12, 31))

        n2 = Nakes.objects.get(no_str='STR-DT-2')
        self.assertEqual(n2.masa_berlaku_str, datetime.date(2027, 8, 15))

        n3 = Nakes.objects.get(no_str='STR-DT-3')
        self.assertIsInstance(n3.masa_berlaku_str, datetime.date)

    def test_upload_case_insensitive_enums(self):
        headers = ['nama', 'profesi', 'unit_kerja', 'no_str', 'masa_berlaku_str', 'no_sip', 'masa_berlaku_sip', 'status_kredensial', 'kewenangan_klinis', 'catatan']
        rows = [
            ['Nakes Lower', 'atlm', 'Lab', 'STR-CI-1', '2026-12-31', 'SIP-CI-1', '2026-12-31', 'dalam proses', 'aktif', ''],
            ['Nakes Mix', 'fisioterapi', 'Rehab', 'STR-CI-2', '2026-12-31', 'SIP-CI-2', '2026-12-31', 'selesai', 'proses', ''],
        ]
        file = self._create_excel_file(rows, headers=headers)
        response = self.client.post(reverse('nakes:upload_bulk'), {'excel_file': file}, follow=True)
        self.assertEqual(response.status_code, 200)

        n1 = Nakes.objects.get(no_str='STR-CI-1')
        self.assertEqual(n1.profesi.nama, 'ATLM')
        self.assertEqual(n1.status_kredensial, 'Dalam Proses')
        self.assertEqual(n1.kewenangan_klinis, 'Aktif')

        n2 = Nakes.objects.get(no_str='STR-CI-2')
        self.assertEqual(n2.profesi.nama, 'Fisioterapis')
        self.assertEqual(n2.status_kredensial, 'Selesai')

    def test_upload_custom_profesi_auto_created(self):
        headers = ['nama', 'profesi', 'unit_kerja', 'no_str', 'masa_berlaku_str', 'no_sip', 'masa_berlaku_sip', 'status_kredensial', 'kewenangan_klinis', 'catatan']
        rows = [
            ['Bidan Sari', 'Bidan Delima', 'KIA', 'STR-BD-1', '2027-12-31', 'SIP-BD-1', '2027-12-31', 'Selesai', 'Aktif', ''],
        ]
        file = self._create_excel_file(rows, headers=headers)
        response = self.client.post(reverse('nakes:upload_bulk'), {'excel_file': file}, follow=True)
        self.assertEqual(response.status_code, 200)

        self.assertTrue(Profesi.objects.filter(nama='Bidan Delima').exists())
        n = Nakes.objects.get(no_str='STR-BD-1')
        self.assertEqual(n.profesi.nama, 'Bidan Delima')

    def test_upload_human_friendly_headers(self):
        headers = ['Nama', 'Profesi', 'Unit Kerja', 'No. STR', 'Masa Berlaku STR', 'No. SIP', 'Masa Berlaku SIP', 'Status Kredensial', 'Kewenangan Klinis', 'Catatan']
        rows = [
            ['Nakes Human', 'Nutrisionis', 'Gizi', 'STR-HF-1', '2026-12-31', 'SIP-HF-1', '2026-12-31', 'Selesai', 'Aktif', 'Test human header'],
        ]
        file = self._create_excel_file(rows, headers=headers)
        response = self.client.post(reverse('nakes:upload_bulk'), {'excel_file': file}, follow=True)
        self.assertEqual(response.status_code, 200)
        self.assertTrue(Nakes.objects.filter(no_str='STR-HF-1').exists())

    def test_sample_rows_ignored(self):
        resp_template = self.client.get(reverse('nakes:download_template'))
        file = SimpleUploadedFile('template.xlsx', resp_template.content, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response = self.client.post(reverse('nakes:upload_bulk'), {'excel_file': file}, follow=True)
        self.assertEqual(response.status_code, 200)
        self.assertFalse(Nakes.objects.filter(no_str='STR-001-2021').exists())

    def test_duplicate_handling(self):
        Nakes.objects.create(
            nama='Existing',
            profesi=self.p_atlm,
            unit_kerja='Lab',
            no_str='STR-DUP-DB',
            masa_berlaku_str=datetime.date(2026, 12, 31),
            no_sip='SIP-DUP-DB',
            masa_berlaku_sip=datetime.date(2026, 12, 31),
            created_by=self.user,
        )
        headers = ['nama', 'profesi', 'unit_kerja', 'no_str', 'masa_berlaku_str', 'no_sip', 'masa_berlaku_sip', 'status_kredensial', 'kewenangan_klinis', 'catatan']
        rows = [
            ['Nakes Dup DB', 'ATLM', 'Lab', 'STR-DUP-DB', '2026-12-31', 'SIP-NEW-1', '2026-12-31', 'Selesai', 'Aktif', ''],
            ['Nakes Unique', 'ATLM', 'Lab', 'STR-UNIQ', '2026-12-31', 'SIP-UNIQ', '2026-12-31', 'Selesai', 'Aktif', ''],
            ['Nakes Dup File', 'ATLM', 'Lab', 'STR-UNIQ', '2026-12-31', 'SIP-UNIQ-2', '2026-12-31', 'Selesai', 'Aktif', ''],
        ]
        file = self._create_excel_file(rows, headers=headers)
        response = self.client.post(reverse('nakes:upload_bulk'), {'excel_file': file}, follow=True)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(Nakes.objects.filter(no_str='STR-UNIQ').count(), 1)


class NakesFormAndExportTest(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username='formuser', password='password123')
        self.client = Client()
        self.client.login(username='formuser', password='password123')

    def test_create_custom_profesi_via_form(self):
        url = reverse('nakes:create')
        res = self.client.post(url, {
            'nama': 'Dr. Sarah Spesialis',
            'profesi': 'Dokter Spesialis Anak',
            'unit_kerja': 'Poli Anak',
            'no_str': 'STR-SARAH-01',
            'masa_berlaku_str': '2028-01-01',
            'no_sip': 'SIP-SARAH-01',
            'masa_berlaku_sip': '2028-01-01',
            'status_kredensial': 'Selesai',
            'kewenangan_klinis': 'Aktif',
            'catatan': '',
        }, follow=True)
        self.assertEqual(res.status_code, 200)

        self.assertTrue(Profesi.objects.filter(nama='Dokter Spesialis Anak').exists())
        nakes = Nakes.objects.get(no_str='STR-SARAH-01')
        self.assertEqual(nakes.profesi.nama, 'Dokter Spesialis Anak')

    def test_export_excel_header_and_row_colors(self):
        today = datetime.date.today()
        p, _ = Profesi.objects.get_or_create(nama='ATLM')

        # 1. Expired (red)
        Nakes.objects.create(
            nama='Nakes Expired',
            profesi=p,
            unit_kerja='Lab',
            no_str='STR-EXP',
            masa_berlaku_str=today - datetime.timedelta(days=10),
            no_sip='SIP-EXP',
            masa_berlaku_sip=today + datetime.timedelta(days=300),
            created_by=self.user,
        )

        # 2. Near expiry <= 180 days (yellow)
        Nakes.objects.create(
            nama='Nakes Warning',
            profesi=p,
            unit_kerja='Lab',
            no_str='STR-WARN',
            masa_berlaku_str=today + datetime.timedelta(days=60),
            no_sip='SIP-WARN',
            masa_berlaku_sip=today + datetime.timedelta(days=300),
            created_by=self.user,
        )

        # 3. Safe > 180 days (green)
        Nakes.objects.create(
            nama='Nakes Safe',
            profesi=p,
            unit_kerja='Lab',
            no_str='STR-SAFE',
            masa_berlaku_str=today + datetime.timedelta(days=400),
            no_sip='SIP-SAFE',
            masa_berlaku_sip=today + datetime.timedelta(days=400),
            created_by=self.user,
        )

        url = reverse('nakes:export')
        res = self.client.get(url)
        self.assertEqual(res.status_code, 200)

        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        ws = wb.active

        header_cell = ws.cell(row=1, column=1)
        self.assertEqual(header_cell.fill.start_color.rgb, '000C7C84')

        color_map = {}
        for row_idx in range(2, ws.max_row + 1):
            nama = ws.cell(row=row_idx, column=1).value
            rgb = ws.cell(row=row_idx, column=1).fill.start_color.rgb
            color_map[nama] = rgb

        self.assertEqual(color_map['Nakes Expired'], '00FFCDD2')
        self.assertEqual(color_map['Nakes Warning'], '00FFF9C4')
        self.assertEqual(color_map['Nakes Safe'], '00C8E6C9')


class DokumenUploadQATest(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username='docuser', password='password123')
        self.client = Client()
        self.client.login(username='docuser', password='password123')
        p, _ = Profesi.objects.get_or_create(nama='ATLM')
        self.nakes = Nakes.objects.create(
            nama='Budi Test',
            profesi=p,
            unit_kerja='Laboratorium',
            no_str='STR-DOC-001',
            masa_berlaku_str=datetime.date(2027, 1, 1),
            no_sip='SIP-DOC-001',
            masa_berlaku_sip=datetime.date(2027, 1, 1),
            created_by=self.user,
        )

    def test_upload_dokumen_nakes_success(self):
        pdf_content = b'%PDF-1.4 test pdf content'
        file = SimpleUploadedFile('str_budi.pdf', pdf_content, content_type='application/pdf')
        url = reverse('nakes:dokumen_upload', args=[self.nakes.pk])
        response = self.client.post(url, {
            'jenis': 'STR',
            'nama_file': 'STR Budi 2027',
            'file': file,
        }, follow=True)
        self.assertEqual(response.status_code, 200)
        self.assertTrue(self.nakes.dokumen.filter(nama_file='STR Budi 2027').exists())

    def test_upload_dokumen_nakes_auto_filename(self):
        pdf_content = b'%PDF-1.4 test pdf content'
        file = SimpleUploadedFile('sip_budi_terbaru.pdf', pdf_content, content_type='application/pdf')
        url = reverse('nakes:dokumen_upload', args=[self.nakes.pk])
        response = self.client.post(url, {
            'jenis': 'SIP',
            'nama_file': '',
            'file': file,
        }, follow=True)
        self.assertEqual(response.status_code, 200)
        self.assertTrue(self.nakes.dokumen.filter(nama_file='sip_budi_terbaru').exists())

    def test_upload_dokumen_invalid_extension(self):
        file = SimpleUploadedFile('malware.exe', b'bad content', content_type='application/octet-stream')
        url = reverse('nakes:dokumen_upload', args=[self.nakes.pk])
        response = self.client.post(url, {
            'jenis': 'STR',
            'file': file,
        }, follow=True)
        self.assertEqual(response.status_code, 200)
        self.assertFalse(self.nakes.dokumen.filter(file__contains='malware').exists())
        messages = [m.message for m in response.context['messages']]
        self.assertTrue(any('tidak didukung' in m for m in messages))

    def test_upload_dokumen_file_too_large(self):
        large_content = b'0' * (16 * 1024 * 1024)
        file = SimpleUploadedFile('huge.pdf', large_content, content_type='application/pdf')
        url = reverse('nakes:dokumen_upload', args=[self.nakes.pk])
        response = self.client.post(url, {
            'jenis': 'STR',
            'file': file,
        }, follow=True)
        self.assertEqual(response.status_code, 200)
        self.assertFalse(self.nakes.dokumen.filter(nama_file='huge').exists())
        messages = [m.message for m in response.context['messages']]
        self.assertTrue(any('maksimal 15 MB' in m for m in messages))

    def test_download_dokumen_nakes(self):
        file = SimpleUploadedFile('cert.pdf', b'%PDF test', content_type='application/pdf')
        doc = DokumenNakes.objects.create(
            nakes=self.nakes,
            jenis='Sertifikat',
            nama_file='Cert 2026',
            file=file,
            uploaded_by=self.user,
        )
        url = reverse('nakes:dokumen_download', args=[doc.pk])
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        response.close()

    def test_dokumen_umum_flow(self):
        file = SimpleUploadedFile('sop_lab.pdf', b'%PDF test sop', content_type='application/pdf')
        url_create = reverse('nakes:dokumen_umum_create')
        res_create = self.client.post(url_create, {
            'judul': 'SOP Laboratorium',
            'kategori': 'SOP',
            'deskripsi': 'Standar lab',
            'file': file,
        })
        self.assertEqual(res_create.status_code, 302)
        self.assertIn('tab=umum', res_create.url)
        doc = DokumenUmum.objects.get(judul='SOP Laboratorium')

        url_dl = reverse('nakes:dokumen_umum_download', args=[doc.pk])
        res_dl = self.client.get(url_dl)
        self.assertEqual(res_dl.status_code, 200)
        res_dl.close()

        url_edit = reverse('nakes:dokumen_umum_edit', args=[doc.pk])
        res_edit = self.client.post(url_edit, {
            'judul': 'SOP Laboratorium Rev 1',
            'kategori': 'SOP',
            'deskripsi': 'Standar lab update',
        })
        self.assertEqual(res_edit.status_code, 302)
        self.assertIn('tab=umum', res_edit.url)
        doc.refresh_from_db()
        self.assertEqual(doc.judul, 'SOP Laboratorium Rev 1')

        url_del = reverse('nakes:dokumen_umum_delete', args=[doc.pk])
        res_del = self.client.post(url_del)
        self.assertEqual(res_del.status_code, 302)
        self.assertIn('tab=umum', res_del.url)
        self.assertFalse(DokumenUmum.objects.filter(pk=doc.pk).exists())


class OPPETest(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username='oppeuser', password='password123')
        self.client = Client()
        self.client.login(username='oppeuser', password='password123')
        p, _ = Profesi.objects.get_or_create(nama='ATLM')
        self.nakes = Nakes.objects.create(
            nama='Dewi OPPE',
            profesi=p,
            unit_kerja='Laboratorium',
            no_str='STR-OPPE-001',
            masa_berlaku_str=datetime.date(2028, 1, 1),
            no_sip='SIP-OPPE-001',
            masa_berlaku_sip=datetime.date(2028, 1, 1),
            created_by=self.user,
        )

    def test_oppe_calculation_and_grade(self):
        # 20 indicators * 85 = 1700 total -> poin 85.00 -> Grade B
        oppe = EvaluasiOPPE.objects.create(
            nakes=self.nakes,
            tahun=2025,
            tanggal_evaluasi=datetime.date(2025, 6, 1),
            skor_perilaku_1=85, skor_perilaku_2=85, skor_perilaku_3=85, skor_perilaku_4=85,
            skor_perilaku_5=85, skor_perilaku_6=85, skor_perilaku_7=85,
            skor_profesional_1=85, skor_profesional_2=85, skor_profesional_3=85, skor_profesional_4=85,
            skor_kinerja_1=85, skor_kinerja_2=85, skor_kinerja_3=85, skor_kinerja_4=85,
            skor_kinerja_5=85, skor_kinerja_6=85, skor_kinerja_7=85, skor_kinerja_8=85, skor_kinerja_9=85,
            penilai_nama='apt. Mariah Ulfah',
            created_by=self.user,
        )
        self.assertEqual(oppe.total_nilai, 1700)
        self.assertEqual(oppe.poin_penilaian, 85.0)
        self.assertEqual(oppe.grade, 'B')
        self.assertEqual(oppe.grade_display, 'Baik')

    def test_oppe_form_renders_all_20_score_inputs(self):
        url_create = reverse('nakes:oppe_create')
        res = self.client.get(url_create)
        self.assertEqual(res.status_code, 200)
        content = res.content.decode('utf-8')
        for i in range(1, 8):
            self.assertIn(f'name="skor_perilaku_{i}"', content)
        for i in range(1, 5):
            self.assertIn(f'name="skor_profesional_{i}"', content)
        for i in range(1, 10):
            self.assertIn(f'name="skor_kinerja_{i}"', content)
        self.assertEqual(content.count('skor-input'), 21)  # 20 inputs + 1 in js querySelector

    def test_oppe_crud_flow(self):
        url_create = reverse('nakes:oppe_create')
        data = {
            'nakes': self.nakes.pk,
            'tahun': 2025,
            'tanggal_evaluasi': '2025-05-10',
            'penilai_nama': 'apt. Mariah Ulfah, S.Farm',
            'penilai_nip': '390.05.11.1',
        }
        for i in range(1, 8):
            data[f'skor_perilaku_{i}'] = 70
        for i in range(1, 5):
            data[f'skor_profesional_{i}'] = 70
        for i in range(1, 10):
            data[f'skor_kinerja_{i}'] = 70

        # Create (Total 1400 -> Poin 70.0 -> Grade C)
        res_create = self.client.post(url_create, data)
        self.assertEqual(res_create.status_code, 302)
        oppe = EvaluasiOPPE.objects.get(nakes=self.nakes, tahun=2025)
        self.assertEqual(oppe.grade, 'C')

        # Detail print view
        url_detail = reverse('nakes:oppe_detail', args=[oppe.pk])
        res_detail = self.client.get(url_detail)
        self.assertEqual(res_detail.status_code, 200)
        self.assertContains(res_detail, 'RS PKU MUHAMMADIYAH GOMBONG')
        self.assertContains(res_detail, 'apt. Mariah Ulfah, S.Farm')

        # Update
        data['skor_kinerja_1'] = 90
        url_update = reverse('nakes:oppe_update', args=[oppe.pk])
        res_update = self.client.post(url_update, data)
        self.assertEqual(res_update.status_code, 302)
        oppe.refresh_from_db()
        self.assertEqual(oppe.total_nilai, 1420)
        self.assertEqual(oppe.poin_penilaian, 71.0)

        # Delete
        url_del = reverse('nakes:oppe_delete', args=[oppe.pk])
        res_del = self.client.post(url_del)
        self.assertEqual(res_del.status_code, 302)
        self.assertFalse(EvaluasiOPPE.objects.filter(pk=oppe.pk).exists())


class MutuKlinisTest(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username='mutuuser', password='password123')
        self.client = Client()
        self.client.login(username='mutuuser', password='password123')

    def test_mutu_crud_and_status(self):
        url_create = reverse('nakes:mutu_create')
        data = {
            'unit_kerja': 'Instalasi Laboratorium',
            'periode_bulan': 5,
            'periode_tahun': 2025,
            'nama_indikator': 'Kepatuhan Pelaporan Nilai Kritis Lab',
            'standar_target': '100.00',
            'capaian': '96.50',
            'analisis': 'Keterlambatan konfirmasi dokter DPJP',
            'rencana_tindak_lanjut': 'Sosialisasi alur pelaporan kritis via WhatsApp',
            'penanggung_jawab': 'dr. Sp.PK',
        }
        res_create = self.client.post(url_create, data)
        self.assertEqual(res_create.status_code, 302)
        mutu = EvaluasiMutuKlinis.objects.get(nama_indikator='Kepatuhan Pelaporan Nilai Kritis Lab')
        self.assertFalse(mutu.is_tercapai)

        # Detail view
        url_detail = reverse('nakes:mutu_detail', args=[mutu.pk])
        res_detail = self.client.get(url_detail)
        self.assertEqual(res_detail.status_code, 200)
        self.assertContains(res_detail, 'Kepatuhan Pelaporan Nilai Kritis Lab')

        # Update capaian to 100% (tercapai)
        data['capaian'] = '100.00'
        url_update = reverse('nakes:mutu_update', args=[mutu.pk])
        res_update = self.client.post(url_update, data)
        self.assertEqual(res_update.status_code, 302)
        mutu.refresh_from_db()
        self.assertTrue(mutu.is_tercapai)

        # Delete
        url_del = reverse('nakes:mutu_delete', args=[mutu.pk])
        res_del = self.client.post(url_del)
        self.assertEqual(res_del.status_code, 302)
        self.assertFalse(EvaluasiMutuKlinis.objects.filter(pk=mutu.pk).exists())


class SubEtikTest(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username='etikuser', password='password123')
        self.client = Client()
        self.client.login(username='etikuser', password='password123')
        p, _ = Profesi.objects.get_or_create(nama='Radiografer')
        self.nakes = Nakes.objects.create(
            nama='Rian Radiologi',
            profesi=p,
            unit_kerja='Radiologi',
            no_str='STR-ETIK-01',
            masa_berlaku_str=datetime.date(2028, 1, 1),
            no_sip='SIP-ETIK-01',
            masa_berlaku_sip=datetime.date(2028, 1, 1),
            created_by=self.user,
        )

    def test_pelanggaran_etik_crud(self):
        # 1. Create
        url_create = reverse('nakes:pelanggaran_create')
        res = self.client.post(url_create, {
            'nakes': self.nakes.pk,
            'tanggal_kejadian': '2025-06-01',
            'tanggal_lapor': '2025-06-02',
            'kategori': 'Sedang',
            'deskripsi': 'Dugaan kelalaian prosedur radiasi',
            'pelapor': 'Kepala Instalasi Radiologi',
            'status': 'Dalam Investigasi',
        })
        self.assertEqual(res.status_code, 302)
        p = PelanggaranEtik.objects.get(nakes=self.nakes)
        self.assertEqual(p.kategori, 'Sedang')
        self.assertEqual(p.status, 'Dalam Investigasi')

        # 2. List
        url_list = reverse('nakes:pelanggaran_list')
        res_list = self.client.get(url_list)
        self.assertEqual(res_list.status_code, 200)
        self.assertContains(res_list, 'Rian Radiologi')

        # 3. Update
        url_update = reverse('nakes:pelanggaran_update', args=[p.pk])
        res_update = self.client.post(url_update, {
            'nakes': self.nakes.pk,
            'tanggal_kejadian': '2025-06-01',
            'tanggal_lapor': '2025-06-02',
            'kategori': 'Sedang',
            'deskripsi': 'Dugaan kelalaian prosedur radiasi update',
            'pelapor': 'Kepala Instalasi Radiologi',
            'status': 'Diteruskan ke Sidang',
        })
        self.assertEqual(res_update.status_code, 302)
        p.refresh_from_db()
        self.assertEqual(p.status, 'Diteruskan ke Sidang')

        # 4. Delete
        url_delete = reverse('nakes:pelanggaran_delete', args=[p.pk])
        res_del = self.client.post(url_delete)
        self.assertEqual(res_del.status_code, 302)
        self.assertFalse(PelanggaranEtik.objects.filter(pk=p.pk).exists())

    def test_sidang_etik_calendar_and_api(self):
        # Create sidang
        sidang = SidangEtik.objects.create(
            nakes=self.nakes,
            judul_sidang='Sidang Etik Kode Etik Radiografer',
            tanggal_sidang=datetime.date(2025, 7, 10),
            waktu_mulai=datetime.time(9, 30),
            waktu_selesai=datetime.time(11, 30),
            tempat='Ruang Komite Medik',
            status='Terjadwal',
            hasil_investigasi='Pemeriksaan bukti rekaman radiologi',
            rekomendasi_pembinaan='Teguran tertulis pertama',
            created_by=self.user,
        )

        # Calendar page view
        url_cal = reverse('nakes:sidang_calendar')
        res_cal = self.client.get(url_cal)
        self.assertEqual(res_cal.status_code, 200)
        self.assertContains(res_cal, 'FullCalendar')

        # JSON API events for calendar
        url_api = reverse('nakes:sidang_events_api')
        res_api = self.client.get(url_api)
        self.assertEqual(res_api.status_code, 200)
        data = res_api.json()
        self.assertEqual(len(data), 1)
        self.assertEqual(data[0]['title'], 'Sidang Etik Kode Etik Radiografer')
        self.assertEqual(data[0]['extendedProps']['nakes'], 'Rian Radiologi')
        self.assertEqual(data[0]['extendedProps']['rekomendasi_pembinaan'], 'Teguran tertulis pertama')

    def test_evaluasi_kinerja_etik_crud(self):
        url_create = reverse('nakes:evaluasi_etik_create')
        res = self.client.post(url_create, {
            'nakes': self.nakes.pk,
            'periode_tahun': 2025,
            'periode_semester': 1,
            'predikat': 'Baik',
            'status_kepatuhan': 'Patuh',
            'catatan_evaluasi': 'Perilaku klinis baik dan sesuai etika profesi',
            'rekomendasi_kelanjutan': 'Direkomendasikan perpanjang SPK',
            'evaluator': 'Sub Komite Etik',
        })
        self.assertEqual(res.status_code, 302)
        ev = EvaluasiKinerjaEtik.objects.get(nakes=self.nakes, periode_tahun=2025, periode_semester=1)
        self.assertEqual(ev.predikat, 'Baik')
        self.assertEqual(ev.status_kepatuhan, 'Patuh')

        # List
        url_list = reverse('nakes:evaluasi_etik_list')
        res_list = self.client.get(url_list)
        self.assertEqual(res_list.status_code, 200)
        self.assertContains(res_list, 'Rian Radiologi')


class SekretariatTest(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username='sekreuser', password='password123')
        self.client = Client()
        self.client.login(username='sekreuser', password='password123')

    def test_agenda_rapat_calendar_and_api(self):
        agenda = AgendaRapat.objects.create(
            judul_rapat='Rapat Pleno Komite KTKL Q3 2025',
            jenis_rapat='Rapat Pleno',
            tanggal_rapat=datetime.date(2025, 8, 15),
            waktu_mulai=datetime.time(9, 0),
            waktu_selesai=datetime.time(11, 0),
            tempat='Ruang Rapat KTKL',
            status='Terjadwal',
            created_by=self.user,
        )
        res_cal = self.client.get(reverse('nakes:agenda_calendar'))
        self.assertEqual(res_cal.status_code, 200)
        self.assertContains(res_cal, 'Agenda Rapat Komite KTKL')

        res_api = self.client.get(reverse('nakes:agenda_events_api'))
        self.assertEqual(res_api.status_code, 200)
        data = res_api.json()
        self.assertEqual(len(data), 1)
        self.assertEqual(data[0]['title'], 'Rapat Pleno Komite KTKL Q3 2025')
        self.assertIn('buat_notulen_url', data[0]['extendedProps'])

        res_create = self.client.post(reverse('nakes:agenda_create'), {
            'judul_rapat': 'Rapat Rutin Bulanan Agustus 2025',
            'jenis_rapat': 'Rutin Bulanan',
            'tanggal_rapat': '2025-08-20',
            'waktu_mulai': '13:00',
            'tempat': 'Ruang Rapat',
            'status': 'Terjadwal',
        })
        self.assertEqual(res_create.status_code, 302)
        self.assertTrue(AgendaRapat.objects.filter(judul_rapat='Rapat Rutin Bulanan Agustus 2025').exists())

    def test_regulasi_crud_and_download_auth(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        pdf_file = SimpleUploadedFile('sk_direktur.pdf', b'%PDF-1.4 dummy content', content_type='application/pdf')
        res_create = self.client.post(reverse('nakes:regulasi_create'), {
            'judul': 'SK Direktur Komite KTKL 2025',
            'nomor_dokumen': '045/SK/DIR/I/2025',
            'kategori': 'SK Direktur',
            'tanggal_terbit': '2025-01-10',
            'tanggal_berlaku': '2025-01-10',
            'status': 'Berlaku',
            'file': pdf_file,
        })
        self.assertEqual(res_create.status_code, 302)
        reg = Regulasi.objects.get(nomor_dokumen='045/SK/DIR/I/2025')
        self.assertEqual(reg.status, 'Berlaku')

        url_dl = reverse('nakes:regulasi_download', args=[reg.pk])
        res_dl = self.client.get(url_dl)
        self.assertEqual(res_dl.status_code, 200)
        res_dl.close()

        anon_client = Client()
        res_anon = anon_client.get(url_dl)
        self.assertEqual(res_anon.status_code, 302)
        self.assertIn('/accounts/', res_anon.url)

        res_list = self.client.get(reverse('nakes:regulasi_list'))
        self.assertEqual(res_list.status_code, 200)
        self.assertContains(res_list, 'SK Direktur Komite KTKL 2025')

    def test_notulen_rapat_crud_and_print(self):
        res_create = self.client.post(reverse('nakes:notulen_create'), {
            'judul': 'Notulen Rapat Pleno KTKL Q3 2025',
            'tanggal': '2025-08-15',
            'waktu_mulai': '09:00',
            'pimpinan_rapat': 'dr. Krisna Widyo, Sp.PK',
            'notulis': 'apt. Mariah Ulfah, S.Farm',
            'agenda_bahasan': 'Evaluasi kredensial semester I 2025',
            'isi_pembahasan': 'Pembahasan hasil OPPE dan tindak lanjut',
            'kesimpulan_keputusan': '1. Nakes A diperpanjang SPK\n2. Nakes B dalam pembinaan',
        })
        self.assertEqual(res_create.status_code, 302)
        notulen = NotulenRapat.objects.get(judul='Notulen Rapat Pleno KTKL Q3 2025')

        url_print = reverse('nakes:notulen_detail', args=[notulen.pk])
        res_print = self.client.get(url_print)
        self.assertEqual(res_print.status_code, 200)
        self.assertContains(res_print, 'RS PKU MUHAMMADIYAH GOMBONG')
        self.assertContains(res_print, 'dr. Krisna Widyo, Sp.PK')

        res_list = self.client.get(reverse('nakes:notulen_list'))
        self.assertEqual(res_list.status_code, 200)
        self.assertContains(res_list, 'Notulen Rapat Pleno KTKL Q3 2025')

