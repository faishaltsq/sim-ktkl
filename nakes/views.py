import re
import datetime
from datetime import date, date as date_type
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.db import transaction
from django.db.models import Q
from django.http import HttpResponse, FileResponse, JsonResponse
from django.urls import reverse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.datetime import from_excel
from openpyxl.worksheet.datavalidation import DataValidation

from .models import (
    Nakes, DokumenNakes, AuditLog, DokumenUmum, Profesi,
    EvaluasiOPPE, EvaluasiMutuKlinis,
    PelanggaranEtik, SidangEtik, EvaluasiKinerjaEtik,
)
from .forms import (
    NakesForm, DokumenForm, NakesSearchForm, DokumenUmumForm,
    EvaluasiOPPEForm, EvaluasiMutuKlinisForm,
    PelanggaranEtikForm, SidangEtikForm, EvaluasiKinerjaEtikForm,
)


HEADER_MAP = {
    'nama': 'nama',
    'namalengkap': 'nama',
    'name': 'nama',
    'profesi': 'profesi',
    'jabatan': 'profesi',
    'profession': 'profesi',
    'unitkerja': 'unit_kerja',
    'unit': 'unit_kerja',
    'departemen': 'unit_kerja',
    'instalasi': 'unit_kerja',
    'nostr': 'no_str',
    'nomorstr': 'no_str',
    'str': 'no_str',
    'masaberlakustr': 'masa_berlaku_str',
    'berlakustr': 'masa_berlaku_str',
    'tglberlakustr': 'masa_berlaku_str',
    'expiredstr': 'masa_berlaku_str',
    'nosip': 'no_sip',
    'nomorsip': 'no_sip',
    'sip': 'no_sip',
    'masaberlakusip': 'masa_berlaku_sip',
    'berlakusip': 'masa_berlaku_sip',
    'tglberlakusip': 'masa_berlaku_sip',
    'expiredsip': 'masa_berlaku_sip',
    'statuskredensial': 'status_kredensial',
    'status': 'status_kredensial',
    'kredensial': 'status_kredensial',
    'kewenanganklinis': 'kewenangan_klinis',
    'kewenangan': 'kewenangan_klinis',
    'catatan': 'catatan',
    'keterangan': 'catatan',
    'note': 'catatan',
}

PROFESI_MAP = {
    'atlm': 'ATLM',
    'atlmteknisilaboratoriummedik': 'ATLM',
    'teknisilaboratoriummedik': 'ATLM',
    'radiografer': 'Radiografer',
    'fisioterapis': 'Fisioterapis',
    'fisioterapi': 'Fisioterapis',
    'nutrisionis': 'Nutrisionis',
    'nutrisi': 'Nutrisionis',
    'gizi': 'Nutrisionis',
    'perekammedis': 'Perekam Medis',
    'rekammedis': 'Perekam Medis',
    'rm': 'Perekam Medis',
    'apoteker': 'Apoteker',
    'sanitarian': 'Sanitarian',
    'lainnya': 'Lainnya',
    'lain': 'Lainnya',
}

STATUS_MAP = {
    'belumpengajuan': 'Belum Pengajuan',
    'belum': 'Belum Pengajuan',
    'belumkredensial': 'Belum Pengajuan',
    'dalamproses': 'Dalam Proses',
    'proses': 'Dalam Proses',
    'sedangproses': 'Dalam Proses',
    'selesai': 'Selesai',
    'sudah': 'Selesai',
}

KEWENANGAN_MAP = {
    'aktif': 'Aktif',
    'evaluasi': 'Evaluasi',
    'proses': 'Proses',
    'tidakaktif': 'Tidak Aktif',
    'nonaktif': 'Tidak Aktif',
}

SAMPLE_STRS = {'STR-001-2021', 'STR-002-2022', 'STR-CONTOH', 'STR-SAMPLE'}


def clean_alphanumeric(val):
    if val is None:
        return ''
    return re.sub(r'[^a-z0-9]', '', str(val).lower())


def parse_date_value(val, label):
    if val is None or str(val).strip() == '':
        return None, f'{label} kosong'
    if isinstance(val, (date_type, datetime.datetime)):
        return val.date() if isinstance(val, datetime.datetime) else val, None
    if isinstance(val, (int, float)):
        try:
            dt = from_excel(val)
            return dt.date() if isinstance(dt, datetime.datetime) else dt, None
        except Exception:
            pass
    s = str(val).strip()
    if len(s) >= 10 and (s[4] == '-' or s[4] == '/'):
        try:
            return date_type.fromisoformat(s[:10].replace('/', '-')), None
        except ValueError:
            pass
    for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y', '%Y/%m/%d', '%d.%m.%Y', '%Y.%m.%d', '%d/%m/%y', '%d-%m-%y'):
        try:
            return datetime.datetime.strptime(s, fmt).date(), None
        except ValueError:
            continue
    return None, f'{label} format salah (gunakan YYYY-MM-DD atau DD/MM/YYYY)'


def log_audit(user, aksi, obj, detail=''):
    AuditLog.objects.create(
        user=user,
        aksi=aksi,
        model_name=obj.__class__.__name__,
        object_id=obj.pk,
        object_repr=str(obj),
        detail=detail,
    )


@login_required
def dashboard(request):
    today = date.today()
    nakes_list = Nakes.objects.all()
    total = nakes_list.count()
    kredensial_selesai = nakes_list.filter(status_kredensial='Selesai').count()
    kewenangan_aktif = nakes_list.filter(kewenangan_klinis='Aktif').count()

    warnings = []
    for n in nakes_list:
        for label, field in [('STR', 'masa_berlaku_str'), ('SIP', 'masa_berlaku_sip')]:
            exp = getattr(n, field)
            sisa = (exp - today).days
            if sisa < 0:
                warnings.append({
                    'level': 'danger',
                    'nakes': n,
                    'doc': label,
                    'tanggal': exp,
                    'sisa': sisa,
                     'msg': f'{n.nama} ({n.profesi.nama}) - {label} sudah kedaluwarsa sejak {exp}',
                })
            elif sisa <= 180:
                warnings.append({
                    'level': 'warning',
                    'nakes': n,
                    'doc': label,
                    'tanggal': exp,
                    'sisa': sisa,
                    'msg': f'{n.nama} ({n.profesi.nama}) - {label} akan habis pada {exp} ({sisa} hari lagi)',
                })

    warnings.sort(key=lambda x: x['sisa'])

    context = {
        'total': total,
        'kredensial_selesai': kredensial_selesai,
        'kewenangan_aktif': kewenangan_aktif,
        'warnings': warnings,
    }
    return render(request, 'nakes/dashboard.html', context)


@login_required
def nakes_list(request):
    form = NakesSearchForm(request.GET)
    qs = Nakes.objects.select_related('profesi').all()
    q = request.GET.get('q', '').strip()
    if q:
        qs = qs.filter(Q(nama__icontains=q) | Q(profesi__nama__icontains=q))
    context = {'nakes_list': qs, 'form': form, 'q': q}
    return render(request, 'nakes/nakes_list.html', context)


@login_required
def nakes_create(request):
    if request.method == 'POST':
        form = NakesForm(request.POST)
        if form.is_valid():
            nakes = form.save(commit=False)
            nakes.created_by = request.user
            nakes.save()
            log_audit(request.user, 'CREATE', nakes)
            messages.success(request, f'Data {nakes.nama} berhasil ditambahkan.')
            return redirect('nakes:detail', pk=nakes.pk)
    else:
        form = NakesForm()
    return render(request, 'nakes/nakes_form.html', {'form': form, 'title': 'Tambah Tenaga Kesehatan'})


@login_required
def nakes_detail(request, pk):
    nakes = get_object_or_404(Nakes, pk=pk)
    dokumen = nakes.dokumen.all()
    dokumen_form = DokumenForm()
    context = {
        'nakes': nakes,
        'dokumen': dokumen,
        'dokumen_form': dokumen_form,
    }
    return render(request, 'nakes/nakes_detail.html', context)


@login_required
def nakes_update(request, pk):
    nakes = get_object_or_404(Nakes, pk=pk)
    if request.method == 'POST':
        form = NakesForm(request.POST, instance=nakes)
        if form.is_valid():
            changed = form.changed_data
            nakes = form.save()
            log_audit(request.user, 'UPDATE', nakes, detail=f'Field diubah: {", ".join(changed)}')
            messages.success(request, f'Data {nakes.nama} berhasil diperbarui.')
            return redirect('nakes:detail', pk=nakes.pk)
    else:
        form = NakesForm(instance=nakes)
    return render(request, 'nakes/nakes_form.html', {'form': form, 'title': f'Edit - {nakes.nama}'})


@login_required
def nakes_delete(request, pk):
    nakes = get_object_or_404(Nakes, pk=pk)
    if request.method == 'POST':
        log_audit(request.user, 'DELETE', nakes)
        nama = nakes.nama
        nakes.delete()
        messages.success(request, f'Data {nama} berhasil dihapus.')
        return redirect('nakes:list')
    return render(request, 'nakes/nakes_confirm_delete.html', {'nakes': nakes})


@login_required
def dokumen_upload(request, nakes_pk):
    nakes = get_object_or_404(Nakes, pk=nakes_pk)
    if request.method == 'POST':
        form = DokumenForm(request.POST, request.FILES)
        if form.is_valid():
            doc = form.save(commit=False)
            doc.nakes = nakes
            doc.uploaded_by = request.user
            if not doc.nama_file and doc.file:
                import os
                doc.nama_file = os.path.splitext(doc.file.name)[0]
            doc.save()
            log_audit(request.user, 'CREATE', doc, detail=f'Upload {doc.jenis} untuk {nakes.nama}')
            messages.success(request, f'Dokumen {doc.jenis} berhasil diupload.')
        else:
            for field, errs in form.errors.items():
                for err in errs:
                    messages.error(request, f'Gagal upload dokumen: {err}')
    return redirect('nakes:detail', pk=nakes_pk)


@login_required
def dokumen_download(request, pk):
    doc = get_object_or_404(DokumenNakes, pk=pk)
    if not doc.file:
        messages.error(request, 'File dokumen tidak ditemukan.')
        return redirect('nakes:detail', pk=doc.nakes.pk)
    try:
        return FileResponse(doc.file.open('rb'))
    except (FileNotFoundError, ValueError):
        messages.error(request, 'File fisik tidak ditemukan pada server penyimpanan.')
        return redirect('nakes:detail', pk=doc.nakes.pk)


@login_required
def dokumen_delete(request, pk):
    doc = get_object_or_404(DokumenNakes, pk=pk)
    nakes_pk = doc.nakes.pk
    if request.method == 'POST':
        log_audit(request.user, 'DELETE', doc)
        try:
            if doc.file:
                doc.file.close()
                doc.file.delete(save=False)
        except Exception:
            pass
        doc.delete()
        messages.success(request, 'Dokumen berhasil dihapus.')
    return redirect('nakes:detail', pk=nakes_pk)


@login_required
def audit_log(request):
    logs = AuditLog.objects.select_related('user').all()[:200]
    return render(request, 'nakes/audit_log.html', {'logs': logs})


@login_required
def export_excel(request):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Data Nakes'

    headers = [
        'Nama', 'Profesi', 'Unit Kerja',
        'No STR', 'Masa Berlaku STR',
        'No SIP', 'Masa Berlaku SIP',
        'Status Kredensial', 'Kewenangan Klinis', 'Catatan',
    ]

    header_fill = PatternFill(start_color='0C7C84', end_color='0C7C84', fill_type='solid')
    header_font = Font(bold=True, color='FFFFFF', size=11)
    center = Alignment(horizontal='center', vertical='center')
    thin = Side(style='thin', color='CCCCCC')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = border

    ws.row_dimensions[1].height = 24

    fill_red = PatternFill(start_color='FFCDD2', end_color='FFCDD2', fill_type='solid')
    fill_yellow = PatternFill(start_color='FFF9C4', end_color='FFF9C4', fill_type='solid')
    fill_green = PatternFill(start_color='C8E6C9', end_color='C8E6C9', fill_type='solid')

    today = date.today()
    for n in Nakes.objects.select_related('profesi').all():
        row_data = [
            n.nama, str(n.profesi), n.unit_kerja,
            n.no_str, n.masa_berlaku_str.isoformat(),
            n.no_sip, n.masa_berlaku_sip.isoformat(),
            n.status_kredensial, n.kewenangan_klinis,
            n.catatan,
        ]
        ws.append(row_data)

        sisa_str = (n.masa_berlaku_str - today).days
        sisa_sip = (n.masa_berlaku_sip - today).days
        worst = min(sisa_str, sisa_sip)

        if worst < 0:
            row_fill = fill_red
        elif worst <= 180:
            row_fill = fill_yellow
        else:
            row_fill = fill_green

        row_idx = ws.max_row
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col)
            cell.fill = row_fill
            cell.border = border

    col_widths = [25, 18, 20, 18, 18, 18, 18, 18, 16, 30]
    for col, width in enumerate(col_widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = width

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="data_nakes.xlsx"'
    wb.save(response)
    return response


@login_required
def download_template(request):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Template Import'

    headers = [
        'nama', 'profesi', 'unit_kerja',
        'no_str', 'masa_berlaku_str',
        'no_sip', 'masa_berlaku_sip',
        'status_kredensial', 'kewenangan_klinis', 'catatan',
    ]

    header_fill = PatternFill(start_color='0C7C84', end_color='0C7C84', fill_type='solid')
    header_font = Font(bold=True, color='FFFFFF')
    center = Alignment(horizontal='center', vertical='center')
    thin = Side(style='thin', color='CCCCCC')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = border

    for row_idx in range(2, 501):
        ws.cell(row=row_idx, column=4).number_format = '@'
        ws.cell(row=row_idx, column=5).number_format = 'yyyy-mm-dd'
        ws.cell(row=row_idx, column=6).number_format = '@'
        ws.cell(row=row_idx, column=7).number_format = 'yyyy-mm-dd'

    q = '"'
    profesi_names = list(Profesi.objects.order_by('nama').values_list('nama', flat=True))
    if not profesi_names:
        profesi_names = ['ATLM', 'Radiografer', 'Fisioterapis', 'Nutrisionis', 'Perekam Medis', 'Apoteker', 'Sanitarian', 'Lainnya']
    dv_profesi = DataValidation(
        type='list',
        formula1=f'{q}{",".join(profesi_names)}{q}',
        allow_blank=True,
    )
    ws.add_data_validation(dv_profesi)
    dv_profesi.add('B2:B500')

    dv_status = DataValidation(
        type='list',
        formula1=f'{q}Belum Pengajuan,Dalam Proses,Selesai{q}',
        allow_blank=True,
    )
    ws.add_data_validation(dv_status)
    dv_status.add('H2:H500')

    dv_kewenangan = DataValidation(
        type='list',
        formula1=f'{q}Aktif,Evaluasi,Proses,Tidak Aktif{q}',
        allow_blank=True,
    )
    ws.add_data_validation(dv_kewenangan)
    dv_kewenangan.add('I2:I500')

    example_rows = [
        ['Budi Santoso', 'ATLM', 'Laboratorium', 'STR-001-2021', '2026-12-31',
         'SIP-001-2021', '2026-12-31', 'Selesai', 'Aktif', ''],
        ['Siti Rahayu', 'Radiografer', 'Radiologi', 'STR-002-2022', '2027-06-30',
         'SIP-002-2022', '2027-06-30', 'Dalam Proses', 'Proses', 'Sedang proses perpanjangan'],
    ]

    note_fill = PatternFill(start_color='F0FAFA', end_color='F0FAFA', fill_type='solid')
    for row_offset, row_data in enumerate(example_rows, start=2):
        for col, val in enumerate(row_data, 1):
            cell = ws.cell(row=row_offset, column=col, value=val)
            cell.fill = note_fill
            cell.border = border

    col_widths = [25, 18, 20, 18, 18, 18, 18, 18, 16, 30]
    for col, width in enumerate(col_widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = width

    ws.row_dimensions[1].height = 22

    info_ws = wb.create_sheet('Panduan')
    info_ws['A1'] = 'PANDUAN IMPORT DATA NAKES'
    info_ws['A1'].font = Font(bold=True, size=13, color='0C7C84')
    panduan = [
        ('', ''),
        ('Kolom', 'Keterangan'),
        ('nama', 'Nama lengkap tenaga kesehatan'),
        ('profesi', 'Nama profesi (sesuai dropdown atau ketik profesi baru)'),
        ('unit_kerja', 'Nama unit/departemen kerja'),
        ('no_str', 'Nomor STR (harus unik)'),
        ('masa_berlaku_str', 'Format: YYYY-MM-DD atau DD/MM/YYYY (contoh: 2026-12-31 atau 31/12/2026)'),
        ('no_sip', 'Nomor SIP (harus unik)'),
        ('masa_berlaku_sip', 'Format: YYYY-MM-DD atau DD/MM/YYYY (contoh: 2026-12-31 atau 31/12/2026)'),
        ('status_kredensial', 'Pilih: Belum Pengajuan | Dalam Proses | Selesai'),
        ('kewenangan_klinis', 'Pilih: Aktif | Evaluasi | Proses | Tidak Aktif'),
        ('catatan', 'Opsional. Catatan tambahan'),
        ('', ''),
        ('PENTING', 'Baris pertama (header) JANGAN diubah'),
        ('PENTING', 'Baris contoh (warna pastel) dapat ditimpa atau dihapus'),
        ('PENTING', 'no_str dan no_sip harus unik; baris duplikat akan dilewati'),
    ]
    for r, (col_a, col_b) in enumerate(panduan, 2):
        info_ws.cell(row=r, column=1, value=col_a)
        info_ws.cell(row=r, column=2, value=col_b)
        if col_a == 'Kolom':
            for c in [1, 2]:
                info_ws.cell(row=r, column=c).font = Font(bold=True)
        if col_a == 'PENTING':
            info_ws.cell(row=r, column=1).font = Font(bold=True, color='C0392B')

    info_ws.column_dimensions['A'].width = 22
    info_ws.column_dimensions['B'].width = 80

    wb.active = 0

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="template_import_nakes.xlsx"'
    wb.save(response)
    return response


@login_required
def upload_bulk(request):
    if request.method == 'POST':
        excel_file = request.FILES.get('excel_file')
        if not excel_file:
            messages.error(request, 'File Excel wajib diupload.')
            return redirect('nakes:upload_bulk')

        try:
            wb = load_workbook(excel_file, data_only=True)
        except Exception:
            messages.error(request, 'File tidak valid. Pastikan format file .xlsx.')
            return redirect('nakes:upload_bulk')

        ws = None
        for sheet_name in wb.sheetnames:
            sheet_candidate = wb[sheet_name]
            first_row = [clean_alphanumeric(c.value) for c in sheet_candidate[1]]
            mapped = {HEADER_MAP[h] for h in first_row if h in HEADER_MAP}
            if {'nama', 'profesi', 'no_str', 'no_sip'}.issubset(mapped):
                ws = sheet_candidate
                break

        if ws is None:
            if 'Template Import' in wb.sheetnames:
                ws = wb['Template Import']
            else:
                ws = wb.active

        first_row_raw = [c.value for c in ws[1]]
        col_map = {}
        for idx, val in enumerate(first_row_raw):
            key = clean_alphanumeric(val)
            field = HEADER_MAP.get(key)
            if field and field not in col_map:
                col_map[field] = idx

        REQUIRED_FIELDS = ['nama', 'profesi', 'unit_kerja', 'no_str', 'masa_berlaku_str', 'no_sip', 'masa_berlaku_sip']
        missing_fields = [f for f in REQUIRED_FIELDS if f not in col_map]
        if missing_fields:
            messages.error(request, f'Header kolom tidak lengkap atau tidak sesuai template. Kolom wajib yang belum ditemukan: {", ".join(missing_fields)}.')
            return redirect('nakes:upload_bulk')

        sukses = 0
        errors = []
        seen_str = set()
        seen_sip = set()

        for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if all(v is None or str(v).strip() == '' for v in row):
                continue

            def get_val(field):
                idx = col_map.get(field)
                if idx is None or idx >= len(row):
                    return ''
                val = row[idx]
                if val is None:
                    return ''
                if isinstance(val, float) and val.is_integer():
                    return str(int(val)).strip()
                return str(val).strip()

            no_str = get_val('no_str')
            no_sip = get_val('no_sip')

            if no_str in SAMPLE_STRS:
                continue

            nama = get_val('nama')
            profesi_raw = get_val('profesi')
            unit_kerja = get_val('unit_kerja')
            status_raw = get_val('status_kredensial')
            kewenangan_raw = get_val('kewenangan_klinis')
            catatan = get_val('catatan')

            idx_masa_str = col_map.get('masa_berlaku_str')
            masa_str_raw = row[idx_masa_str] if idx_masa_str is not None and idx_masa_str < len(row) else None

            idx_masa_sip = col_map.get('masa_berlaku_sip')
            masa_sip_raw = row[idx_masa_sip] if idx_masa_sip is not None and idx_masa_sip < len(row) else None

            row_errors = []

            if not nama:
                row_errors.append('nama kosong')

            if not profesi_raw:
                row_errors.append('profesi kosong')
                profesi_obj = None
            else:
                profesi_name = PROFESI_MAP.get(clean_alphanumeric(profesi_raw), profesi_raw.strip())
                profesi_obj, _ = Profesi.objects.get_or_create(nama=profesi_name)

            if not unit_kerja:
                row_errors.append('unit_kerja kosong')

            if not no_str:
                row_errors.append('no_str kosong')

            if not no_sip:
                row_errors.append('no_sip kosong')

            status_norm = STATUS_MAP.get(clean_alphanumeric(status_raw), 'Belum Pengajuan') if status_raw else 'Belum Pengajuan'
            if status_raw and not STATUS_MAP.get(clean_alphanumeric(status_raw)):
                row_errors.append(f'status_kredensial tidak valid: "{status_raw}"')

            kewenangan_norm = KEWENANGAN_MAP.get(clean_alphanumeric(kewenangan_raw), 'Proses') if kewenangan_raw else 'Proses'
            if kewenangan_raw and not KEWENANGAN_MAP.get(clean_alphanumeric(kewenangan_raw)):
                row_errors.append(f'kewenangan_klinis tidak valid: "{kewenangan_raw}"')

            masa_str, err_str = parse_date_value(masa_str_raw, 'masa_berlaku_str')
            if err_str:
                row_errors.append(err_str)

            masa_sip, err_sip = parse_date_value(masa_sip_raw, 'masa_berlaku_sip')
            if err_sip:
                row_errors.append(err_sip)

            if row_errors:
                errors.append(f'Baris {row_num}: {", ".join(row_errors)}')
                continue

            if no_str in seen_str:
                errors.append(f'Baris {row_num}: no_str "{no_str}" duplikat di dalam file, dilewati')
                continue
            seen_str.add(no_str)

            if no_sip in seen_sip:
                errors.append(f'Baris {row_num}: no_sip "{no_sip}" duplikat di dalam file, dilewati')
                continue
            seen_sip.add(no_sip)

            if Nakes.objects.filter(no_str=no_str).exists():
                errors.append(f'Baris {row_num}: no_str "{no_str}" sudah ada di sistem, dilewati')
                continue
            if Nakes.objects.filter(no_sip=no_sip).exists():
                errors.append(f'Baris {row_num}: no_sip "{no_sip}" sudah ada di sistem, dilewati')
                continue

            try:
                with transaction.atomic():
                    nakes = Nakes.objects.create(
                        nama=nama,
                        profesi=profesi_obj,
                        unit_kerja=unit_kerja,
                        no_str=no_str,
                        masa_berlaku_str=masa_str,
                        no_sip=no_sip,
                        masa_berlaku_sip=masa_sip,
                        status_kredensial=status_norm,
                        kewenangan_klinis=kewenangan_norm,
                        catatan=catatan,
                        created_by=request.user,
                    )
                    log_audit(request.user, 'CREATE', nakes, detail='Import bulk Excel')
                    sukses += 1
            except Exception as ex:
                errors.append(f'Baris {row_num}: Gagal simpan ke database ({str(ex)})')

        if sukses:
            messages.success(request, f'{sukses} data berhasil diimport.')
        if errors:
            if len(errors) > 5:
                for e in errors[:5]:
                    messages.warning(request, e)
                messages.warning(request, f'... dan {len(errors) - 5} baris lainnya bermasalah atau dilewati.')
            else:
                for e in errors:
                    messages.warning(request, e)
        if not sukses and not errors:
            messages.info(request, 'Tidak ada data yang diproses. File mungkin kosong.')

        return redirect('nakes:upload_bulk')

    return render(request, 'nakes/upload_bulk.html')


@login_required
def dokumen_hub(request):
    tab = request.GET.get('tab', 'nakes')
    q = request.GET.get('q', '').strip()
    kategori = request.GET.get('kategori', '')

    dokumen_nakes = DokumenNakes.objects.select_related('nakes', 'uploaded_by').all()
    if q:
        dokumen_nakes = dokumen_nakes.filter(
            Q(nama_file__icontains=q) | Q(nakes__nama__icontains=q)
        )

    dokumen_umum = DokumenUmum.objects.select_related('uploaded_by').all()
    if q and tab == 'umum':
        dokumen_umum = dokumen_umum.filter(
            Q(judul__icontains=q) | Q(deskripsi__icontains=q)
        )
    if kategori and tab == 'umum':
        dokumen_umum = dokumen_umum.filter(kategori=kategori)

    context = {
        'tab': tab,
        'q': q,
        'kategori': kategori,
        'dokumen_nakes': dokumen_nakes,
        'dokumen_umum': dokumen_umum,
        'kategori_choices': DokumenUmum.KATEGORI_CHOICES,
    }
    return render(request, 'nakes/dokumen_hub.html', context)


@login_required
def dokumen_umum_create(request):
    if request.method == 'POST':
        form = DokumenUmumForm(request.POST, request.FILES)
        if form.is_valid():
            doc = form.save(commit=False)
            doc.uploaded_by = request.user
            doc.save()
            log_audit(request.user, 'CREATE', doc, detail=f'Upload dokumen umum: {doc.judul}')
            messages.success(request, f'Dokumen "{doc.judul}" berhasil diupload.')
            return redirect(f"{reverse('nakes:dokumen_hub')}?tab=umum")
        else:
            for field, errs in form.errors.items():
                for err in errs:
                    messages.error(request, f'Gagal upload dokumen: {err}')
    else:
        form = DokumenUmumForm()
    return render(request, 'nakes/dokumen_umum_form.html', {'form': form, 'title': 'Upload Dokumen Umum'})


@login_required
def dokumen_umum_edit(request, pk):
    doc = get_object_or_404(DokumenUmum, pk=pk)
    if request.method == 'POST':
        form = DokumenUmumForm(request.POST, request.FILES, instance=doc)
        if form.is_valid():
            changed = form.changed_data
            doc = form.save()
            log_audit(request.user, 'UPDATE', doc, detail=f'Field diubah: {", ".join(changed)}')
            messages.success(request, f'Dokumen "{doc.judul}" berhasil diperbarui.')
            return redirect(f"{reverse('nakes:dokumen_hub')}?tab=umum")
        else:
            for field, errs in form.errors.items():
                for err in errs:
                    messages.error(request, f'Gagal memperbarui dokumen: {err}')
    else:
        form = DokumenUmumForm(instance=doc)
    return render(request, 'nakes/dokumen_umum_form.html', {'form': form, 'title': f'Edit - {doc.judul}'})


@login_required
def dokumen_umum_download(request, pk):
    doc = get_object_or_404(DokumenUmum, pk=pk)
    if not doc.file:
        messages.error(request, 'File dokumen tidak ditemukan.')
        return redirect(f"{reverse('nakes:dokumen_hub')}?tab=umum")
    try:
        return FileResponse(doc.file.open('rb'))
    except (FileNotFoundError, ValueError):
        messages.error(request, 'File fisik tidak ditemukan pada server penyimpanan.')
        return redirect(f"{reverse('nakes:dokumen_hub')}?tab=umum")


@login_required
def dokumen_umum_delete(request, pk):
    doc = get_object_or_404(DokumenUmum, pk=pk)
    if request.method == 'POST':
        log_audit(request.user, 'DELETE', doc)
        try:
            if doc.file:
                doc.file.close()
                doc.file.delete(save=False)
        except Exception:
            pass
        doc.delete()
        messages.success(request, 'Dokumen berhasil dihapus.')
    return redirect(f"{reverse('nakes:dokumen_hub')}?tab=umum")


# ==========================================
# SUB MUTU & PENGEMBANGAN PROFESI: OPPE
# ==========================================

@login_required
def oppe_list(request):
    qs = EvaluasiOPPE.objects.select_related('nakes', 'nakes__profesi').all()
    q = request.GET.get('q', '').strip()
    tahun = request.GET.get('tahun', '').strip()
    grade = request.GET.get('grade', '').strip()

    if q:
        qs = qs.filter(Q(nakes__nama__icontains=q) | Q(nakes__unit_kerja__icontains=q))
    if tahun:
        qs = qs.filter(tahun=tahun)
    if grade:
        qs = qs.filter(grade=grade)

    years = EvaluasiOPPE.objects.values_list('tahun', flat=True).distinct().order_by('-tahun')

    context = {
        'oppe_list': qs,
        'q': q,
        'tahun': tahun,
        'grade': grade,
        'years': years,
    }
    return render(request, 'nakes/oppe_list.html', context)


@login_required
def oppe_create(request):
    if request.method == 'POST':
        form = EvaluasiOPPEForm(request.POST)
        if form.is_valid():
            oppe = form.save(commit=False)
            oppe.created_by = request.user
            oppe.save()
            log_audit(request.user, 'CREATE', oppe, detail=f'Penilaian OPPE {oppe.nakes.nama} ({oppe.tahun}) Grade: {oppe.grade}')
            messages.success(request, f'Evaluasi OPPE {oppe.nakes.nama} tahun {oppe.tahun} berhasil disimpan.')
            return redirect('nakes:oppe_detail', pk=oppe.pk)
    else:
        initial = {
            'tahun': date.today().year,
            'tanggal_evaluasi': date.today(),
        }
        form = EvaluasiOPPEForm(initial=initial)

    return render(request, 'nakes/oppe_form.html', {
        'form': form,
        'title': 'Tambah Evaluasi OPPE',
    })


@login_required
def oppe_detail(request, pk):
    oppe = get_object_or_404(EvaluasiOPPE.objects.select_related('nakes', 'nakes__profesi'), pk=pk)
    return render(request, 'nakes/oppe_print.html', {'oppe': oppe})


@login_required
def oppe_update(request, pk):
    oppe = get_object_or_404(EvaluasiOPPE, pk=pk)
    if request.method == 'POST':
        form = EvaluasiOPPEForm(request.POST, instance=oppe)
        if form.is_valid():
            oppe = form.save()
            log_audit(request.user, 'UPDATE', oppe, detail=f'Update OPPE {oppe.nakes.nama} ({oppe.tahun}) Grade: {oppe.grade}')
            messages.success(request, f'Evaluasi OPPE {oppe.nakes.nama} tahun {oppe.tahun} berhasil diperbarui.')
            return redirect('nakes:oppe_detail', pk=oppe.pk)
    else:
        form = EvaluasiOPPEForm(instance=oppe)

    return render(request, 'nakes/oppe_form.html', {
        'form': form,
        'title': f'Edit Evaluasi OPPE - {oppe.nakes.nama}',
        'oppe': oppe,
    })


@login_required
def oppe_delete(request, pk):
    oppe = get_object_or_404(EvaluasiOPPE, pk=pk)
    if request.method == 'POST':
        log_audit(request.user, 'DELETE', oppe, detail=f'Hapus OPPE {oppe.nakes.nama} ({oppe.tahun})')
        oppe.delete()
        messages.success(request, 'Evaluasi OPPE berhasil dihapus.')
    return redirect('nakes:oppe_list')


# ==========================================
# SUB MUTU: EVALUASI MUTU LAYANAN KLINIS
# ==========================================

@login_required
def mutu_list(request):
    qs = EvaluasiMutuKlinis.objects.all()
    q = request.GET.get('q', '').strip()
    unit = request.GET.get('unit', '').strip()
    tahun = request.GET.get('tahun', '').strip()
    bulan = request.GET.get('bulan', '').strip()

    if q:
        qs = qs.filter(Q(nama_indikator__icontains=q) | Q(unit_kerja__icontains=q))
    if unit:
        qs = qs.filter(unit_kerja=unit)
    if tahun:
        qs = qs.filter(periode_tahun=tahun)
    if bulan:
        qs = qs.filter(periode_bulan=bulan)

    units = EvaluasiMutuKlinis.objects.values_list('unit_kerja', flat=True).distinct().order_by('unit_kerja')
    years = EvaluasiMutuKlinis.objects.values_list('periode_tahun', flat=True).distinct().order_by('-periode_tahun')

    total_count = qs.count()
    tercapai_count = sum(1 for m in qs if m.is_tercapai)
    belum_tercapai_count = total_count - tercapai_count

    context = {
        'mutu_list': qs,
        'q': q,
        'unit': unit,
        'tahun': tahun,
        'bulan': bulan,
        'units': units,
        'years': years,
        'bulan_choices': EvaluasiMutuKlinis.BULAN_CHOICES,
        'total_count': total_count,
        'tercapai_count': tercapai_count,
        'belum_tercapai_count': belum_tercapai_count,
    }
    return render(request, 'nakes/mutu_list.html', context)


@login_required
def mutu_create(request):
    if request.method == 'POST':
        form = EvaluasiMutuKlinisForm(request.POST)
        if form.is_valid():
            mutu = form.save(commit=False)
            mutu.created_by = request.user
            mutu.save()
            log_audit(request.user, 'CREATE', mutu, detail=f'Indikator Mutu: {mutu.nama_indikator} ({mutu.unit_kerja})')
            messages.success(request, f'Indikator Mutu "{mutu.nama_indikator}" berhasil ditambahkan.')
            return redirect('nakes:mutu_detail', pk=mutu.pk)
    else:
        initial = {
            'periode_tahun': date.today().year,
            'periode_bulan': date.today().month,
            'standar_target': 100.0,
        }
        form = EvaluasiMutuKlinisForm(initial=initial)

    return render(request, 'nakes/mutu_form.html', {'form': form, 'title': 'Tambah Indikator Mutu Klinis'})


@login_required
def mutu_detail(request, pk):
    mutu = get_object_or_404(EvaluasiMutuKlinis, pk=pk)
    return render(request, 'nakes/mutu_detail.html', {'mutu': mutu})


@login_required
def mutu_update(request, pk):
    mutu = get_object_or_404(EvaluasiMutuKlinis, pk=pk)
    if request.method == 'POST':
        form = EvaluasiMutuKlinisForm(request.POST, instance=mutu)
        if form.is_valid():
            mutu = form.save()
            log_audit(request.user, 'UPDATE', mutu, detail=f'Update Indikator Mutu: {mutu.nama_indikator}')
            messages.success(request, f'Indikator Mutu "{mutu.nama_indikator}" berhasil diperbarui.')
            return redirect('nakes:mutu_detail', pk=mutu.pk)
    else:
        form = EvaluasiMutuKlinisForm(instance=mutu)

    return render(request, 'nakes/mutu_form.html', {'form': form, 'title': f'Edit - {mutu.nama_indikator}'})


@login_required
def mutu_delete(request, pk):
    mutu = get_object_or_404(EvaluasiMutuKlinis, pk=pk)
    if request.method == 'POST':
        log_audit(request.user, 'DELETE', mutu, detail=f'Hapus Mutu: {mutu.nama_indikator}')
        mutu.delete()
        messages.success(request, 'Indikator mutu klinis berhasil dihapus.')
    return redirect('nakes:mutu_list')


# ==========================================
# SUB ETIK: PENCATATAN PELANGGARAN ETIK
# ==========================================

@login_required
def pelanggaran_list(request):
    qs = PelanggaranEtik.objects.select_related('nakes', 'nakes__profesi').all()
    q = request.GET.get('q', '').strip()
    kategori = request.GET.get('kategori', '').strip()
    status = request.GET.get('status', '').strip()

    if q:
        qs = qs.filter(Q(nakes__nama__icontains=q) | Q(nakes__unit_kerja__icontains=q) | Q(deskripsi__icontains=q))
    if kategori:
        qs = qs.filter(kategori=kategori)
    if status:
        qs = qs.filter(status=status)

    context = {
        'pelanggaran_list': qs,
        'q': q,
        'kategori': kategori,
        'status': status,
    }
    return render(request, 'nakes/etik/pelanggaran_list.html', context)


@login_required
def pelanggaran_create(request):
    if request.method == 'POST':
        form = PelanggaranEtikForm(request.POST)
        if form.is_valid():
            obj = form.save(commit=False)
            obj.created_by = request.user
            obj.save()
            log_audit(request.user, 'CREATE', obj, detail=f'Pelanggaran etik: {obj.nakes.nama}')
            messages.success(request, f'Pelanggaran etik {obj.nakes.nama} berhasil dicatat.')
            return redirect('nakes:pelanggaran_list')
    else:
        form = PelanggaranEtikForm(initial={'tanggal_lapor': date.today()})
    return render(request, 'nakes/etik/pelanggaran_form.html', {'form': form, 'title': 'Catat Pelanggaran Etik'})


@login_required
def pelanggaran_update(request, pk):
    obj = get_object_or_404(PelanggaranEtik, pk=pk)
    if request.method == 'POST':
        form = PelanggaranEtikForm(request.POST, instance=obj)
        if form.is_valid():
            form.save()
            log_audit(request.user, 'UPDATE', obj, detail=f'Update pelanggaran etik: {obj.nakes.nama}')
            messages.success(request, 'Data pelanggaran berhasil diperbarui.')
            return redirect('nakes:pelanggaran_list')
    else:
        form = PelanggaranEtikForm(instance=obj)
    return render(request, 'nakes/etik/pelanggaran_form.html', {'form': form, 'title': f'Edit - {obj.nakes.nama}'})


@login_required
def pelanggaran_delete(request, pk):
    obj = get_object_or_404(PelanggaranEtik, pk=pk)
    if request.method == 'POST':
        log_audit(request.user, 'DELETE', obj, detail=f'Hapus pelanggaran etik: {obj.nakes.nama}')
        obj.delete()
        messages.success(request, 'Data pelanggaran berhasil dihapus.')
    return redirect('nakes:pelanggaran_list')


# ==========================================
# SUB ETIK: SIDANG & REKOMENDASI PEMBINAAN
# ==========================================

@login_required
def sidang_calendar(request):
    return render(request, 'nakes/etik/sidang_calendar.html')


@login_required
def sidang_events_api(request):
    qs = SidangEtik.objects.select_related('nakes', 'nakes__profesi').all()
    color_map = {
        'Terjadwal': '#0C7C84',
        'Selesai': '#059669',
        'Ditunda': '#D97706',
        'Dibatalkan': '#DC2626',
    }
    events = []
    for s in qs:
        events.append({
            'id': s.pk,
            'title': s.judul_sidang,
            'start': f'{s.tanggal_sidang}T{s.waktu_mulai.strftime("%H:%M")}',
            'end': f'{s.tanggal_sidang}T{s.waktu_selesai.strftime("%H:%M")}' if s.waktu_selesai else None,
            'color': color_map.get(s.status, '#64748B'),
            'extendedProps': {
                'pk': s.pk,
                'nakes': s.nakes.nama,
                'profesi': str(s.nakes.profesi),
                'unit': s.nakes.unit_kerja,
                'tempat': s.tempat,
                'status': s.status,
                'hasil_investigasi': s.hasil_investigasi,
                'rekomendasi_pembinaan': s.rekomendasi_pembinaan,
                'tindak_lanjut': s.tindak_lanjut,
                'edit_url': reverse('nakes:sidang_update', args=[s.pk]),
            },
        })
    return JsonResponse(events, safe=False)


@login_required
def sidang_create(request):
    if request.method == 'POST':
        form = SidangEtikForm(request.POST)
        if form.is_valid():
            obj = form.save(commit=False)
            obj.created_by = request.user
            obj.save()
            log_audit(request.user, 'CREATE', obj, detail=f'Sidang etik: {obj.judul_sidang}')
            messages.success(request, f'Sidang etik "{obj.judul_sidang}" berhasil dijadwalkan.')
            return redirect('nakes:sidang_calendar')
    else:
        form = SidangEtikForm(initial={'tanggal_sidang': date.today(), 'waktu_mulai': '09:00'})
    return render(request, 'nakes/etik/sidang_form.html', {'form': form, 'title': 'Jadwalkan Sidang Etik'})


@login_required
def sidang_update(request, pk):
    obj = get_object_or_404(SidangEtik, pk=pk)
    if request.method == 'POST':
        form = SidangEtikForm(request.POST, instance=obj)
        if form.is_valid():
            form.save()
            log_audit(request.user, 'UPDATE', obj, detail=f'Update sidang: {obj.judul_sidang}')
            messages.success(request, 'Data sidang berhasil diperbarui.')
            return redirect('nakes:sidang_calendar')
    else:
        form = SidangEtikForm(instance=obj)
    return render(request, 'nakes/etik/sidang_form.html', {'form': form, 'title': f'Edit - {obj.judul_sidang}'})


@login_required
def sidang_delete(request, pk):
    obj = get_object_or_404(SidangEtik, pk=pk)
    if request.method == 'POST':
        log_audit(request.user, 'DELETE', obj, detail=f'Hapus sidang: {obj.judul_sidang}')
        obj.delete()
        messages.success(request, 'Data sidang berhasil dihapus.')
    return redirect('nakes:sidang_calendar')


# ==========================================
# SUB ETIK: RIWAYAT EVALUASI KINERJA ETIK
# ==========================================

@login_required
def evaluasi_etik_list(request):
    qs = EvaluasiKinerjaEtik.objects.select_related('nakes', 'nakes__profesi').all()
    q = request.GET.get('q', '').strip()
    tahun = request.GET.get('tahun', '').strip()
    predikat = request.GET.get('predikat', '').strip()

    if q:
        qs = qs.filter(Q(nakes__nama__icontains=q) | Q(nakes__unit_kerja__icontains=q))
    if tahun:
        qs = qs.filter(periode_tahun=tahun)
    if predikat:
        qs = qs.filter(predikat=predikat)

    years = EvaluasiKinerjaEtik.objects.values_list('periode_tahun', flat=True).distinct().order_by('-periode_tahun')

    context = {
        'evaluasi_list': qs,
        'q': q,
        'tahun': tahun,
        'predikat': predikat,
        'years': years,
    }
    return render(request, 'nakes/etik/evaluasi_etik_list.html', context)


@login_required
def evaluasi_etik_create(request):
    if request.method == 'POST':
        form = EvaluasiKinerjaEtikForm(request.POST)
        if form.is_valid():
            obj = form.save(commit=False)
            obj.created_by = request.user
            obj.save()
            log_audit(request.user, 'CREATE', obj, detail=f'Evaluasi etik: {obj.nakes.nama} ({obj.periode_tahun} Sem.{obj.periode_semester})')
            messages.success(request, f'Evaluasi kinerja etik {obj.nakes.nama} berhasil disimpan.')
            return redirect('nakes:evaluasi_etik_list')
    else:
        form = EvaluasiKinerjaEtikForm(initial={'periode_tahun': date.today().year})
    return render(request, 'nakes/etik/evaluasi_etik_form.html', {'form': form, 'title': 'Tambah Evaluasi Kinerja Etik'})


@login_required
def evaluasi_etik_update(request, pk):
    obj = get_object_or_404(EvaluasiKinerjaEtik, pk=pk)
    if request.method == 'POST':
        form = EvaluasiKinerjaEtikForm(request.POST, instance=obj)
        if form.is_valid():
            form.save()
            log_audit(request.user, 'UPDATE', obj, detail=f'Update evaluasi etik: {obj.nakes.nama}')
            messages.success(request, 'Data evaluasi kinerja etik berhasil diperbarui.')
            return redirect('nakes:evaluasi_etik_list')
    else:
        form = EvaluasiKinerjaEtikForm(instance=obj)
    return render(request, 'nakes/etik/evaluasi_etik_form.html', {'form': form, 'title': f'Edit - {obj.nakes.nama}'})


@login_required
def evaluasi_etik_delete(request, pk):
    obj = get_object_or_404(EvaluasiKinerjaEtik, pk=pk)
    if request.method == 'POST':
        log_audit(request.user, 'DELETE', obj, detail=f'Hapus evaluasi etik: {obj.nakes.nama}')
        obj.delete()
        messages.success(request, 'Data evaluasi kinerja etik berhasil dihapus.')
    return redirect('nakes:evaluasi_etik_list')
